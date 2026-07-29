from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches

md = Path("BillBanaoPay_Website_Marketing_Brief.md").read_text(encoding="utf-8")
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

for line in md.splitlines():
    s = line.rstrip()
    if not s:
        continue
    if s.startswith("# "):
        doc.add_heading(s[2:], 0)
    elif s.startswith("## "):
        doc.add_heading(s[3:], 1)
    elif s.startswith("### "):
        doc.add_heading(s[4:], 2)
    elif s.startswith("- ") or s.startswith("* "):
        doc.add_paragraph(s[2:], style="List Bullet")
    elif s.startswith("|"):
        p = doc.add_paragraph(s)
        for r in p.runs:
            r.font.size = Pt(9)
    elif s.startswith(">"):
        p = doc.add_paragraph(s[1:].strip())
        if p.runs:
            p.runs[0].italic = True
    elif s.startswith("---") or s.startswith("```"):
        continue
    else:
        doc.add_paragraph(s.replace("**", ""))

doc.add_heading("Screenshot pack (visual)", 1)
shots = Path("screenshots")
if shots.exists():
    for f in sorted(shots.glob("*.jpg")):
        doc.add_paragraph(f.stem)
        try:
            doc.add_picture(str(f), width=Inches(5.2))
        except Exception:
            doc.add_paragraph("(see " + f.name + ")")

out = Path("BillBanaoPay Website Marketing Brief - OmniDEL.docx")
doc.save(out)
print(out.resolve(), out.stat().st_size)
